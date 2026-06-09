from __future__ import annotations

import re
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional


_IMAGE_EMBED_RE = re.compile(r"!\[\[(.*?)\]\]")
MAX_TEXT_LINES = 10


def compile_latex(tex_path: Path, max_runs: int = 5) -> None:
    cwd = tex_path.parent
    tex_name = tex_path.name
    rerun = True
    runs = 0
    while rerun and runs < max_runs:
        runs += 1
        subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", tex_name],
            cwd=cwd,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            check=False,
        )
        log_file = cwd / tex_path.with_suffix(".log").name
        rerun = False
        if log_file.exists():
            log_text = log_file.read_text(encoding="utf-8", errors="ignore")
            if (
                "Rerun to get cross-references right" in log_text
                or "Label(s) may have changed" in log_text
                or "undefined references" in log_text
            ):
                rerun = True


def _parse_bool(value: Optional[str], default: bool = False) -> bool:
    if value is None:
        return default
    cleaned = str(value).strip().strip('"').strip("'").lower()
    if cleaned in {"true", "yes", "1", "on"}:
        return True
    if cleaned in {"false", "no", "0", "off"}:
        return False
    return default


def _pandoc_convert(markdown: str) -> str:
    try:
        import pypandoc  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "Missing optional dependency 'pypandoc'. Install pandoc and pypandoc to enable Beamer export."
        ) from exc
    try:
        return pypandoc.convert_text(markdown, "latex", format="md", extra_args=["--wrap=preserve"])
    except OSError as exc:
        raise RuntimeError("Pandoc not found. Install pandoc and ensure it is on PATH.") from exc


def convert_obsidian_images(md_text: str) -> str:
    def repl(match: re.Match) -> str:
        inner = match.group(1).strip()
        parts = inner.split("|")
        img = parts[0].strip()

        background = False
        width = None
        height = None

        for part in parts[1:]:
            part = part.strip()
            if re.match(r"^\d+x\d+$", part):
                width, height = part.split("x")
            elif part.lower().startswith(("inil", "inll")):
                background = True

        size_opts: List[str] = []
        if width and height:
            size_opts.append(f"width={width}pt")
            size_opts.append(f"height={height}pt")
            size_opts.append("keepaspectratio")
        size_str = ",".join(size_opts)

        if background:
            if size_opts:
                return (
                    "\\usebackgroundtemplate{%\n"
                    f"  \\includegraphics[{size_str}]{{{img}}}\n"
                    "}"
                )
            return (
                "\\usebackgroundtemplate{%\n"
                f"  \\includegraphics[width=\\paperwidth,height=\\paperheight]{{{img}}}\n"
                "}"
            )
        if size_opts:
            return f"\\includegraphics[{size_str}]{{{img}}}"
        return f"\\includegraphics[width=0.8\\textwidth]{{{img}}}"

    return re.sub(r"!\[\[(.*?)\]\]", repl, md_text)


def convert_markdown_table(md_table: str) -> str:
    lines = [line.strip() for line in md_table.strip().splitlines() if line.strip()]
    header = [h.strip() for h in lines[0].strip("|").split("|")]
    align_line = lines[1].strip("|").split("|")

    aligns = []
    for al in align_line:
        al = al.strip()
        if al.startswith(":") and al.endswith(":"):
            aligns.append("c")
        elif al.startswith(":"):
            aligns.append("l")
        elif al.endswith(":"):
            aligns.append("r")
        else:
            aligns.append("l")

    body = []
    for line in lines[2:]:
        row = [c.strip() for c in line.strip("|").split("|")]
        body.append(row)

    col_format_parts = []
    for al in aligns:
        if al == "l":
            col_format_parts.append(">{\\raggedright\\arraybackslash}X")
        elif al == "c":
            col_format_parts.append(">{\\centering\\arraybackslash}X")
        elif al == "r":
            col_format_parts.append(">{\\raggedleft\\arraybackslash}X")
    col_format = "|".join(col_format_parts)

    def esc(cell: str) -> str:
        cell = (
            cell.replace("&", "\\&")
            .replace("%", "\\%")
            .replace("_", "\\_")
            .replace("#", "\\#")
            .replace("{", "\\{")
            .replace("}", "\\}")
        )
        cell = re.sub(r"(?<!\\)\$", r"\\$", cell)
        return cell

    table_latex = "\\resizebox{\\textwidth}{!}{%\n"
    table_latex += "\\begin{tabularx}{\\textwidth}{" + col_format + "}\n"
    table_latex += " \\hline\n "
    table_latex += " & ".join(esc(c) for c in header) + " \\\\\n \\hline\n"
    for row in body:
        table_latex += " " + " & ".join(esc(c) for c in row) + " \\\\\n"
    table_latex += " \\hline\n\\end{tabularx}\n}%"
    return table_latex


def convert_md_body(md_text: str) -> str:
    md_text = convert_obsidian_images(md_text)

    lines = md_text.splitlines()
    out_lines: List[str] = []
    in_table = False
    table_buffer: List[str] = []

    for line in lines:
        if re.match(r"^\s*\|.*\|\s*$", line):
            in_table = True
            table_buffer.append(line)
            continue
        if in_table:
            out_lines.append(convert_markdown_table("\n".join(table_buffer)))
            table_buffer = []
            in_table = False
        out_lines.append(line)

    if in_table:
        out_lines.append(convert_markdown_table("\n".join(table_buffer)))

    final_parts: List[str] = []
    for part in out_lines:
        if (
            part.strip().startswith("\\resizebox")
            or part.strip().startswith("\\includegraphics")
            or part.strip().startswith("\\usebackgroundtemplate")
        ):
            final_parts.append(part)
        elif part.strip():
            final_parts.append(_pandoc_convert(part))
        else:
            final_parts.append("")

    return "\n".join(final_parts)


def _extract_first_image(md_text: str) -> tuple[Optional[str], str]:
    match = _IMAGE_EMBED_RE.search(md_text)
    if not match:
        return None, md_text
    inner = match.group(1).strip()
    img = inner.split("|")[0].strip()
    if not img.replace("\\", "/").lower().startswith("images/"):
        return None, md_text
    cleaned = _IMAGE_EMBED_RE.sub("", md_text, count=1)
    cleaned = "\n".join(line for line in cleaned.splitlines() if line.strip())
    return img or None, cleaned.strip()


def _split_markdown_lines(md_text: str, max_lines: int = MAX_TEXT_LINES) -> List[str]:
    lines = [line.strip() for line in md_text.splitlines() if line.strip()]
    if not lines:
        return [""]
    if len(lines) <= max_lines:
        return ["\n".join(lines)]
    chunks: List[str] = []
    for idx in range(0, len(lines), max_lines):
        chunk_lines = lines[idx : idx + max_lines]
        chunks.append("\n".join(chunk_lines))
    return chunks


def _strip_markdown_formatting(text: str) -> str:
    cleaned = text.strip()
    cleaned = re.sub(r"!\[\[(.*?)\]\]", "", cleaned)
    cleaned = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", cleaned)
    cleaned = re.sub(r"`{1,3}(.+?)`{1,3}", r"\1", cleaned)
    cleaned = re.sub(r"(\*\*|__)(.*?)\1", r"\2", cleaned)
    cleaned = re.sub(r"(?<!\*)\*(?!\*)(.*?)\*(?<!\*)", r"\1", cleaned)
    cleaned = re.sub(r"(?<!_)_(?!_)(.*?)_(?<!_)", r"\1", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip()


def _markdown_lines_for_pptx(md_text: str) -> List[tuple[str, int]]:
    lines: List[tuple[str, int]] = []
    for raw_line in md_text.splitlines():
        if not raw_line.strip():
            continue
        line = raw_line.rstrip()
        indent = len(line) - len(line.lstrip(" "))
        level = min(indent // 2, 4)
        stripped = line.strip()
        numbered = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        bulleted = re.match(r"^[-*+]\s+(.*)$", stripped)
        if numbered:
            display = f"{numbered.group(1)}. {_strip_markdown_formatting(numbered.group(2))}"
        elif bulleted:
            display = f"• {_strip_markdown_formatting(bulleted.group(1))}"
        else:
            display = _strip_markdown_formatting(stripped)
        if display:
            lines.append((display, level))
    return lines


@dataclass
class SlideFrame:
    index: int
    title: str
    body: str
    raw: str


def split_presentation_markdown(md_text: str) -> List[SlideFrame]:
    frames_raw = md_text.split("\n---\n")
    slides: List[SlideFrame] = []
    for idx, frame_text in enumerate(frames_raw[1:], start=1):
        title = ""
        match = re.search(r"^\s*#\s+(.+)$", frame_text, flags=re.MULTILINE)
        if match:
            title = match.group(1).strip()
        body = re.sub(r"^\s*#\s+.+$", "", frame_text, flags=re.MULTILINE)
        body = body.replace("<!-- _class: title -->", "").strip()
        slides.append(SlideFrame(index=idx, title=title, body=body, raw=frame_text))
    return slides


def parse_slide_ranges(slide_ranges_str: str) -> set[int]:
    slides_to_remove: set[int] = set()
    if not slide_ranges_str.strip():
        return slides_to_remove
    for part in slide_ranges_str.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = map(int, part.split("-"))
            slides_to_remove.update(range(start, end + 1))
        else:
            slides_to_remove.add(int(part))
    return slides_to_remove


def md_to_beamer_tex(input_file: Path) -> Path:
    input_path = Path(input_file)
    file_stem = input_path.stem

    content = input_path.read_text(encoding="utf-8")
    frames_raw = content.split("\n---\n")

    params: Dict[str, str] = {}
    for line in frames_raw[0].splitlines():
        match = re.match(r"^(\w+)\s*:\s*(.+)$", line.strip())
        if match:
            key, value = match.groups()
            params[key.strip().lower()] = value.strip()

    theme = params.get("theme", "Madrid")
    if theme.lower() == "beamer":
        theme = "Madrid"
    paginate = _parse_bool(params.get("paginate"), default=False)
    outline = _parse_bool(params.get("outline"), default=True)
    author = params.get("author", "")
    header = params.get("header", "")
    footer_text = params.get("footer", "")

    presentation_title = file_stem
    beamer_frames: List[str] = []

    for frame_text in frames_raw[1:]:
        if "<!-- _class: title -->" in frame_text:
            match = re.search(r"^\s*#\s+(.+)", frame_text, flags=re.MULTILINE)
            if match:
                presentation_title = match.group(1).strip()
            frame_body_md = re.sub(r"^\s*#\s+.+", "", frame_text, flags=re.MULTILINE)
            frame_body_md = frame_body_md.replace("<!-- _class: title -->", "")
            latex_body = convert_md_body(frame_body_md)

            author_line = f"\\author{{{author}}}" if author else "\\author{}"
            title_frame = (
                "\\begin{frame}[plain]\n"
                f"\\title{{{presentation_title}}}\n"
                f"{author_line}\n"
                "\\maketitle\n"
                f"{latex_body}\n"
                "\\end{frame}\n"
            )
            beamer_frames.append(title_frame)
            if outline:
                beamer_frames.append(
                    "\\begin{frame}[allowframebreaks]{Outline}\n\\tableofcontents\n\\end{frame}\n"
                )
            continue

        match = re.search(r"^\s*#\s+(.+)", frame_text, flags=re.MULTILINE)
        frame_title = ""
        section_line = ""
        if match:
            frame_title = match.group(1).strip()
            section_line = "\\section{" + frame_title + "}\n"
            frame_body_md = re.sub(r"^\s*#\s+.+", "", frame_text, flags=re.MULTILINE)
        else:
            frame_body_md = frame_text

        image_path, cleaned_body = _extract_first_image(frame_body_md)
        body_chunks = [cleaned_body]
        if image_path:
            body_chunks = _split_markdown_lines(cleaned_body, MAX_TEXT_LINES)

        for chunk_idx, chunk_md in enumerate(body_chunks):
            chunk_md = re.sub(r"^\s*##\s+(.+)$", r"\\subsection{\1}", chunk_md, flags=re.MULTILINE)
            latex_body = convert_md_body(chunk_md)

            allow_breaks = ""
            if not image_path and "\\begin{tabularx}" in latex_body:
                allow_breaks = "[allowframebreaks]"

            chunk_title = frame_title
            chunk_section = section_line if chunk_idx == 0 else ""
            if chunk_idx > 0:
                chunk_title = f"{frame_title} (cont.)" if frame_title else ""

            if image_path:
                image_latex = (
                    "\\begin{columns}[T,onlytextwidth]\n"
                    "\\begin{column}{0.58\\textwidth}\n"
                    "\\scriptsize\n"
                    f"{latex_body}\n"
                    "\\end{column}\n"
                    "\\begin{column}{0.42\\textwidth}\n"
                    "\\centering\n"
                    f"\\includegraphics[width=\\linewidth,height=0.8\\textheight,keepaspectratio]{{{image_path}}}\n"
                    "\\end{column}\n"
                    "\\end{columns}\n"
                )
                frame_code = (
                    chunk_section
                    + f"\\begin{{frame}}{{{chunk_title}}}\n"
                    + image_latex
                    + "\\end{frame}\n"
                )
            else:
                frame_code = (
                    chunk_section
                    + f"\\begin{{frame}}{allow_breaks}{{{chunk_title}}}\n"
                    + "\\scriptsize\n"
                    + latex_body
                    + "\n\\end{frame}\n"
                )
            beamer_frames.append(frame_code)

    paginate_code = ""
    if paginate:
        paginate_code = r"""
\setbeamertemplate{footline}[frame number]
"""

    custom_footer = ""
    if header or footer_text:
        right_code = footer_text
        pages = ""
        if paginate:
            pages = " \\hspace{1em} \\insertframenumber{} / \\inserttotalframenumber"
        custom_footer = (
            r"""
\setbeamertemplate{footline}{
  \leavevmode%
  \hbox{%
  \begin{beamercolorbox}[wd=.5\paperwidth,ht=2.25ex,dp=1ex,leftskip=1em]{author in head/foot}%
    """
            + header
            + r"""%
  \end{beamercolorbox}%
  \begin{beamercolorbox}[wd=.4\paperwidth,ht=2.25ex,dp=1ex,rightskip=1em plus1fil]{title in head/foot}%
    \centering """
            + right_code
            + r"""%
  \end{beamercolorbox}}%
  \begin{beamercolorbox}[wd=.1\paperwidth,ht=2.25ex,dp=1ex,rightskip=1em plus2fil]{pages in head/foot}%
    \raggedleft """
            + pages
            + r"""%
  \end{beamercolorbox}}%
  \vskip0pt%
}
"""
        )

    preamble = (
        "\\documentclass[aspectratio=169]{beamer}\n"
        "\\setbeamersize{text margin left=0pt,text margin right=0pt}\n"
        "\\usetheme{" + theme + "}\n"
        + paginate_code
        + custom_footer
        + "\n"
    )
    if author:
        preamble += "\\author{" + author + "}\n"
    preamble += "\\title{" + presentation_title + "}\n\\date{}\n"
    preamble += r"""
\usepackage[utf8]{inputenc}
\usepackage{ulem}
\usepackage{graphicx}
\usepackage{amsmath}
\usepackage{hyperref}
\usepackage{tabularx}
\usepackage{array}
\renewcommand{\arraystretch}{1.2}
"""

    tex_content = preamble + "\\begin{document}\n" + "".join(beamer_frames) + "\\end{document}\n"

    tex_file = input_path.with_suffix(".tex")
    tex_file.write_text(tex_content, encoding="utf-8")
    return tex_file


def md_to_pptx(input_file: Path) -> Path:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN  # type: ignore
        from pptx.util import Inches, Pt  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "Missing optional dependency 'python-pptx'. Install it to enable PPTX export."
        ) from exc

    input_path = Path(input_file)
    content = input_path.read_text(encoding="utf-8", errors="ignore")
    slides = split_presentation_markdown(content)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]
    blank_layout = prs.slide_layouts[6]

    def _set_text_frame(text_frame: Any, lines: List[tuple[str, int]], *, font_size: int) -> None:
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        first = True
        for text, level in lines:
            paragraph = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
            first = False
            paragraph.text = text
            paragraph.level = level
            paragraph.font.size = Pt(font_size)
            paragraph.alignment = PP_ALIGN.LEFT

    for slide in slides:
        image_path, cleaned_body = _extract_first_image(slide.body)
        body_lines = _markdown_lines_for_pptx(cleaned_body)
        is_title_slide = "<!-- _class: title -->" in slide.raw

        if is_title_slide:
            ppt_slide = prs.slides.add_slide(title_layout)
            ppt_slide.shapes.title.text = _strip_markdown_formatting(slide.title or input_path.stem)
            subtitle = ppt_slide.placeholders[1]
            subtitle.text = "\n".join(text for text, _ in body_lines) if body_lines else ""
            continue

        if image_path:
            ppt_slide = prs.slides.add_slide(blank_layout)
            title_box = ppt_slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.6))
            title_frame = title_box.text_frame
            title_frame.text = _strip_markdown_formatting(slide.title)
            title_frame.paragraphs[0].font.size = Pt(24)
            title_frame.paragraphs[0].font.bold = True

            body_box = ppt_slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(6.2), Inches(5.9))
            _set_text_frame(body_box.text_frame, body_lines or [("", 0)], font_size=16)

            resolved_image = (input_path.parent / image_path).resolve()
            if resolved_image.exists():
                ppt_slide.shapes.add_picture(
                    str(resolved_image),
                    Inches(7.0),
                    Inches(1.1),
                    width=Inches(5.8),
                    height=Inches(5.2),
                )
            continue

        ppt_slide = prs.slides.add_slide(content_layout)
        ppt_slide.shapes.title.text = _strip_markdown_formatting(slide.title)
        body_placeholder = ppt_slide.placeholders[1]
        _set_text_frame(body_placeholder.text_frame, body_lines or [("", 0)], font_size=18)

    pptx_path = input_path.with_suffix(".pptx")
    prs.save(str(pptx_path))
    return pptx_path
